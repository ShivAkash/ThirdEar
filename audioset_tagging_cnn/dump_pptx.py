from pptx import Presentation
import sys

def dump_pptx(filename):
    try:
        prs = Presentation(filename)
        for i, slide in enumerate(prs.slides):
            print(f"--- Slide {i+1} ---")
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    print(shape.text)
    except Exception as e:
        print(f"Error: {e}")

if __name__ == "__main__":
    dump_pptx(sys.argv[1])
